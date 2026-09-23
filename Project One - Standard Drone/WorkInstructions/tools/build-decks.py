#!/usr/bin/env python3
"""Build a PowerPoint deck from each WI-*.md slide deck in WorkInstructions/.

The markdown is the source of truth. Each `## Slide N: title` section becomes
one slide; the photo named in its `<!-- photo: ... -->` comment or `<img src>`
tag is placed if the file exists, otherwise the slide shows a labelled slot.

    python tools/build-decks.py            # every WI-*.md -> WI-*.pptx
    python tools/build-decks.py WI-01      # just one

Requires: pip install python-pptx pillow
"""

from __future__ import annotations

import math
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROUTING = Path(__file__).resolve().parent.parent

# ---------------------------------------------------------------- palette --
# Workshop palette: carbon-fibre charcoal dominates, solder-amber is the one
# accent, slate for anything that should recede.
CHARCOAL = RGBColor(0x1F, 0x23, 0x26)
INK = RGBColor(0x2B, 0x30, 0x34)
SLATE = RGBColor(0x6B, 0x75, 0x7C)
AMBER = RGBColor(0xE8, 0xA3, 0x3D)
AMBER_TINT = RGBColor(0xFD, 0xF3, 0xE3)
CARD = RGBColor(0xF2, 0xF3, 0xF4)
FRAME = RGBColor(0xE3, 0xE5, 0xE7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Calibri"
SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)


# ---------------------------------------------------------------- parsing --
@dataclass
class Block:
    kind: str                    # label | list | callout | check | text | next
    text: str = ""
    items: list[str] = field(default_factory=list)
    ordered: bool = False


@dataclass
class Slide:
    number: int
    title: str
    meta: str = ""               # "Op 30 · 30 min"
    photos: list[str] = field(default_factory=list)  # relative paths, e.g. images/wi-01/07-x.jpg
    blocks: list[Block] = field(default_factory=list)


@dataclass
class Deck:
    wi: str                      # "WI-01"
    title: str                   # "Bench setup, frame assembly, motor mounting"
    facts: list[tuple[str, str]] # (Status, "Complete — 2026-09-20"), ...
    slides: list[Slide]


INLINE_STRIP = [
    (re.compile(r"\[([^\]]+)\]\([^)]+\)"), r"\1"),   # [text](url) -> text
    (re.compile(r"~~([^~]+)~~"), r"\1"),
    (re.compile(r"`([^`]+)`"), r"\1"),
    (re.compile(r"(?<!\*)\*(?!\*)([^*]+)\*(?!\*)"), r"\1"),  # *italic*
]


def plain(s: str) -> str:
    """Strip links/italics/strike; keep **bold** markers for the renderer."""
    for rx, rep in INLINE_STRIP:
        s = rx.sub(rep, s)
    return re.sub(r"\s+", " ", s).strip()


def paragraphs(body: str) -> list[str]:
    out, cur = [], []
    for line in body.splitlines():
        if line.strip():
            cur.append(line.rstrip())
        elif cur:
            out.append("\n".join(cur))
            cur = []
    if cur:
        out.append("\n".join(cur))
    return out


LIST_RX = re.compile(r"^\s*(?:(\d+)\.|[-*])\s+(.*)$")


def parse_list(par: str) -> Block | None:
    items, ordered, any_item = [], False, False
    for line in par.splitlines():
        m = LIST_RX.match(line)
        if m:
            any_item = True
            ordered = ordered or bool(m.group(1))
            items.append(m.group(2))
        elif items and line.startswith((" ", "\t")):
            items[-1] += " " + line.strip()
        else:
            return None
    return Block("list", items=[plain(i) for i in items], ordered=ordered) if any_item else None


def parse_slide_body(body: str) -> tuple[str, list[str], list[Block]]:
    meta, photos, blocks = "", [], []
    for par in paragraphs(body):
        s = par.strip()
        if s.startswith("<sub>") or re.fullmatch(r"-{3,}", s):
            continue
        # One or more photo slots / images in a row — a paragraph can hold several.
        found = re.findall(r"<!--\s*photo:\s*(\S+)\s*-->", s) + re.findall(r'<img[^>]*src="([^"]+)"', s)
        if found and re.fullmatch(r"(?:\s*(?:<!--[^>]*-->|<img[^>]*>)\s*)+", s):
            photos.extend(found)
            continue
        if re.fullmatch(r"\*[^*].*[^*]\*", s) and "\n" not in s and not blocks:
            meta = plain(s)
            continue
        lst = parse_list(par)
        if lst:
            # "**Check**" label directly before a list -> that list is the check
            if blocks and blocks[-1].kind == "label" and blocks[-1].text.lower().rstrip(":") == "check":
                blocks.pop()
                lst.kind = "check"
            blocks.append(lst)
            continue
        flat = plain(s)
        m = re.match(r"\*\*Check:\*\*\s*(.*)", flat, re.S)
        if m:
            blocks.append(Block("check", text=m.group(1)))
            continue
        m = re.match(r"\*\*Next:\*\*\s*(.*)", flat, re.S)
        if m:
            blocks.append(Block("next", text=m.group(1)))
            continue
        m = re.fullmatch(r"\*\*([^*]+)\*\*(?:\s*[—-]\s*(.*))?", flat)
        if m and len(m.group(1)) <= 60:
            blocks.append(Block("label", text=m.group(1) + (" — " + m.group(2) if m.group(2) else "")))
            continue
        if re.fullmatch(r"\*\*.+\*\*", flat):
            blocks.append(Block("callout", text=flat.strip("*")))
            continue
        blocks.append(Block("text", text=flat))
    return meta, photos, blocks


def parse_deck(md: Path) -> Deck:
    text = md.read_text(encoding="utf-8")
    head = re.search(r"^# (\S+)\s+—\s+(.*)$", text, re.M)
    if not head:
        sys.exit(f"{md.name}: no '# WI-NN — title' heading")
    wi, title = head.group(1), plain(head.group(2))

    facts = []
    for m in re.finditer(r"^\| \*\*([^*]+)\*\* \| (.*) \|$", text, re.M):
        facts.append((m.group(1), plain(m.group(2)).replace("**", "")))

    slides = []
    parts = re.split(r"^## Slide (\d+): (.*)$", text, flags=re.M)
    for i in range(1, len(parts), 3):
        number, stitle, body = int(parts[i]), plain(parts[i + 1]), parts[i + 2]
        meta, photos, blocks = parse_slide_body(body)
        slides.append(Slide(number, stitle, meta, photos, blocks))
    if not slides:
        sys.exit(f"{md.name}: no '## Slide N:' sections")
    return Deck(wi, title, facts, slides)


# -------------------------------------------------------------- rendering --
def add_text(slide, x, y, w, h, *, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, margin=0):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(margin)
    p = tf.paragraphs[0]
    p.alignment = align
    tb._defaults = dict(size=size, color=color, bold=bold)
    return tb


def runs(par, text, size, color, bold=False, italic=False):
    """Write `text` into paragraph `par`, honouring **bold** spans."""
    for i, chunk in enumerate(re.split(r"(\*\*[^*]+\*\*)", text)):
        if not chunk:
            continue
        r = par.add_run()
        if chunk.startswith("**") and chunk.endswith("**"):
            r.text, r.font.bold = chunk[2:-2], True
        else:
            r.text, r.font.bold = chunk, bold
        r.font.size, r.font.name, r.font.italic = Pt(size), FONT, italic
        r.font.color.rgb = color


def add_rect(slide, x, y, w, h, fill, rounded=True, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    if rounded:
        shp.adjustments[0] = 0.06
    shp.shadow.inherit = False
    return shp


def number_badge(slide, x, y, n, d=Inches(0.62), fill=AMBER, color=CHARCOAL, size=20):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
    c.fill.solid()
    c.fill.fore_color.rgb = fill
    c.line.fill.background()
    c.shadow.inherit = False
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    runs(p, str(n), size, color, bold=True)
    return c


def est_lines(text: str, chars_per_line: int) -> int:
    return max(1, math.ceil(len(text.replace("**", "")) / chars_per_line))


def footer(slide, deck: Deck, n: int, total: int, dark=False):
    col = SLATE if not dark else RGBColor(0x9A, 0xA3, 0xA9)
    tb = add_text(slide, Inches(0.6), Inches(6.95), Inches(9), Inches(0.3), anchor=MSO_ANCHOR.MIDDLE)
    runs(tb.text_frame.paragraphs[0], f"{deck.wi} · {deck.title}", 10, col)
    tb = add_text(slide, Inches(8.2), Inches(6.95), Inches(4.5), Inches(0.3), align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    runs(tb.text_frame.paragraphs[0], f"{n} / {total}", 10, col)


def set_background(slide, color):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = color


def title_slide(prs, deck: Deck):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(s, CHARCOAL)
    tb = add_text(s, Inches(0.8), Inches(1.2), Inches(11), Inches(0.5))
    runs(tb.text_frame.paragraphs[0], f"PROJECT ONE  ·  ROUTING  ·  {deck.wi.upper()}", 14, AMBER, bold=True)
    tb = add_text(s, Inches(0.8), Inches(1.8), Inches(11.5), Inches(1.6), anchor=MSO_ANCHOR.TOP)
    runs(tb.text_frame.paragraphs[0], deck.title, 40, WHITE, bold=True)
    y = Inches(4.0)
    for k, v in deck.facts:
        tb = add_text(s, Inches(0.8), y, Inches(1.6), Inches(0.5))
        runs(tb.text_frame.paragraphs[0], k.upper(), 11, AMBER, bold=True)
        tb = add_text(s, Inches(2.5), y, Inches(10), Inches(0.5))
        runs(tb.text_frame.paragraphs[0], v, 14, RGBColor(0xE6, 0xE8, 0xEA))
        y += Inches(0.18) + Pt(14) * max(1, est_lines(v, 120)) * 1.3
    footer(s, deck, 0, 0, dark=True)
    s.shapes[-1].text_frame.paragraphs[0].runs[0].text = "One slide per step · one photo · one check"


def deck_slide(prs, deck: Deck):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(s, WHITE)
    tb = add_text(s, Inches(0.6), Inches(0.5), Inches(11), Inches(0.9))
    runs(tb.text_frame.paragraphs[0], "Deck", 32, CHARCOAL, bold=True)
    n = len(deck.slides)
    per_col = math.ceil(n / 2)
    row_h = min(Inches(0.5), int(Inches(4.9) / per_col))
    for i, sl in enumerate(deck.slides):
        col, row = divmod(i, per_col)
        x = Inches(0.6) + col * Inches(6.3)
        y = Inches(1.7) + row * row_h
        number_badge(s, x, y, sl.number, d=Inches(0.4), size=12)
        tb = add_text(s, x + Inches(0.55), y, Inches(4.4), Inches(0.4), anchor=MSO_ANCHOR.MIDDLE)
        runs(tb.text_frame.paragraphs[0], sl.title, 15, INK)
        if sl.meta:
            tb = add_text(s, x + Inches(4.5), y, Inches(1.0), Inches(0.4), anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
            runs(tb.text_frame.paragraphs[0], sl.meta.split("·")[0].strip(), 11, SLATE)
    footer(s, deck, 0, 0)
    s.shapes[-1].text_frame.paragraphs[0].runs[0].text = f"{n} slides"


def fit_picture(slide, path: Path, x, y, w, h):
    """Place the image centred inside the box, preserving aspect ratio."""
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = int(iw * scale), int(ih * scale)
    slide.shapes.add_picture(str(path), x + (w - pw) // 2, y + (h - ph) // 2, pw, ph)


def place_photos(slide, sl: Slide, x, y, w, h, dark=False):
    """One frame; 1–3 photos side by side inside it. Missing files show a labelled slot."""
    add_rect(slide, x, y, w, h, INK if dark else CARD, line=SLATE if dark else FRAME)
    slots = sl.photos or [""]
    pad, gutter = Inches(0.12), Inches(0.1)
    n = len(slots)
    cw = (w - 2 * pad - gutter * (n - 1)) // n
    for i, rel in enumerate(slots):
        cx = x + pad + i * (cw + gutter)
        path = ROUTING / rel if rel else None
        if path and path.exists():
            fit_picture(slide, path, cx, y + pad, cw, h - 2 * pad)
            continue
        tb = add_text(slide, cx, y, cw, h, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
        p = tb.text_frame.paragraphs[0]
        runs(p, "PHOTO", 13, SLATE, bold=True)
        if rel:
            p2 = tb.text_frame.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            runs(p2, rel.rsplit("/", 1)[-1], 10, SLATE)


def content_slide(prs, deck: Deck, sl: Slide, total: int, end=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    dark = end
    set_background(s, CHARCOAL if dark else WHITE)
    text_col = WHITE if dark else INK
    muted = RGBColor(0xB8, 0xBF, 0xC4) if dark else SLATE

    number_badge(s, Inches(0.6), Inches(0.55), sl.number)
    tb = add_text(s, Inches(1.45), Inches(0.42), Inches(11.2), Inches(0.9), anchor=MSO_ANCHOR.MIDDLE)
    runs(tb.text_frame.paragraphs[0], sl.title, 30, WHITE if dark else CHARCOAL, bold=True)
    if sl.meta:
        tb = add_text(s, Inches(1.45), Inches(1.3), Inches(8), Inches(0.35))
        runs(tb.text_frame.paragraphs[0], sl.meta.upper(), 11, AMBER if dark else SLATE, bold=True)

    # Two columns: words left, photos right. Slides with several photos give
    # the photo frame more of the width (phone photos are portrait).
    many = len(sl.photos) > 1
    lx, ly, lw = Inches(0.6), Inches(1.85), Inches(4.7 if many else 6.1)
    px, py, pw, ph = lx + lw + Inches(0.4), Inches(1.85), Inches(12.75) - (lx + lw + Inches(0.4)), Inches(4.75)
    place_photos(s, sl, px, py, pw, ph, dark=dark)

    # Pick a body size that fits the column: estimate total lines first.
    cpl_at = {16: 40, 14: 46, 12: 54} if many else {16: 52, 14: 60, 12: 70}
    def total_lines(size):
        cpl, n = cpl_at[size], 0
        for b in sl.blocks:
            if b.kind == "list" or b.kind == "check" and b.items:
                n += sum(est_lines(i, cpl - 4) for i in b.items) + 1.5
            else:
                n += est_lines(b.text, cpl) + 1.5
        return n
    size = next((sz for sz in (16, 14, 12) if total_lines(sz) * sz * 1.3 <= 4.9 * 72), 12)
    cpl = cpl_at[size]
    line_h = Pt(size * 1.3)
    gap = Inches(0.18)

    def card_height(b: Block) -> int:
        lines = (sum(est_lines(i, cpl - 6) for i in b.items) + 1) if b.items else est_lines("CHECK  " + b.text, cpl - 4)
        return int(lines * line_h + Inches(0.32))

    # The CHECK card is anchored to the bottom of the column on every slide,
    # so it sits in the same place deck-wide and can never collide with the
    # footer. Everything else flows down from the top.
    y = ly
    column_bottom = Inches(6.75)
    checks = [b for b in sl.blocks if b.kind == "check"]
    if checks:
        b = checks[-1]
        h = card_height(b)
        cy = column_bottom - h
        add_rect(s, lx, cy, lw, h, INK if dark else CARD)
        tb = add_text(s, lx + Inches(0.2), cy + Inches(0.15), lw - Inches(0.4), h - Inches(0.3), margin=0)
        p = tb.text_frame.paragraphs[0]
        runs(p, "CHECK" + ("  " if not b.items else ""), size - 3, AMBER, bold=True)
        if b.items:
            for it in b.items:
                p = tb.text_frame.add_paragraph()
                p.space_before = Pt(3)
                runs(p, "✓  " + it, size - 1, text_col)
        else:
            runs(p, b.text, size - 1, text_col)

    for b in sl.blocks:
        if b.kind == "check":
            continue
        if b.kind == "label":
            tb = add_text(s, lx, y, lw, line_h)
            runs(tb.text_frame.paragraphs[0], b.text.upper(), size - 3, AMBER, bold=True)
            y += line_h + Inches(0.04)
        elif b.kind == "callout":
            h = card_height(b)
            add_rect(s, lx, y, lw, h, AMBER_TINT if not dark else INK)
            tb = add_text(s, lx + Inches(0.2), y + Inches(0.15), lw - Inches(0.4), h - Inches(0.3), margin=0)
            p = tb.text_frame.paragraphs[0]
            runs(p, "CAUTION  ", size - 3, AMBER, bold=True)
            runs(p, b.text, size - 1, text_col, bold=True)
            y += h + gap
        elif b.kind == "list":
            tb = add_text(s, lx, y, lw, Inches(4))
            first = True
            for i, it in enumerate(b.items, 1):
                p = tb.text_frame.paragraphs[0] if first else tb.text_frame.add_paragraph()
                first = False
                p.space_after = Pt(size * 0.45)
                prefix = f"{i}.  " if b.ordered else "•  "
                runs(p, prefix, size, AMBER, bold=True)
                runs(p, it, size, text_col)
            y += int(sum(est_lines(i, cpl - 4) for i in b.items) * line_h + len(b.items) * Pt(size * 0.45)) + gap
        elif b.kind == "next":
            tb = add_text(s, lx, y, lw, line_h * 2)
            p = tb.text_frame.paragraphs[0]
            runs(p, "NEXT  ", size - 3, AMBER, bold=True)
            runs(p, b.text, size, text_col, bold=True)
            y += line_h * est_lines(b.text, cpl) + gap
        else:
            tb = add_text(s, lx, y, lw, line_h * est_lines(b.text, cpl) + Inches(0.1))
            runs(tb.text_frame.paragraphs[0], b.text, size, muted if not dark else text_col)
            y += line_h * est_lines(b.text, cpl) + gap

    footer(s, deck, sl.number, total, dark=dark)


def build(md: Path) -> Path:
    deck = parse_deck(md)
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    title_slide(prs, deck)
    deck_slide(prs, deck)
    total = len(deck.slides)
    for sl in deck.slides:
        content_slide(prs, deck, sl, total, end=(sl is deck.slides[-1]))
    out = md.with_suffix(".pptx")
    prs.save(out)
    return out


def main(argv: list[str]) -> None:
    targets = sorted(ROUTING.glob("WI-*.md"))
    if argv:
        targets = [t for t in targets if any(t.name.startswith(a) for a in argv)]
    if not targets:
        sys.exit("no WI-*.md decks matched")
    for md in targets:
        out = build(md)
        deck = parse_deck(md)
        wanted = [p for sl in deck.slides for p in sl.photos]
        placed = [p for p in wanted if (ROUTING / p).exists()]
        print(f"{out.name}: {len(deck.slides)} slides, {len(placed)}/{len(wanted)} photos placed")
        for p in wanted:
            if p not in placed:
                print(f"  missing: {p}")


if __name__ == "__main__":
    main(sys.argv[1:])
