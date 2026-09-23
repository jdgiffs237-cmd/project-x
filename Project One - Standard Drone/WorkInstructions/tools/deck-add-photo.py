#!/usr/bin/env python3
"""Put a photo into an empty slot of the master slide deck, in place.

    python tools/deck-add-photo.py --list
    python tools/deck-add-photo.py 5 images/wi-01/02-iron-at-temperature.jpg
    python tools/deck-add-photo.py 5 images/wi-01/02b-second-angle.jpg   # again: side by side

The deck (Project-One-Build-Walkthrough.pptx) is edited in Google Slides and
is the master — this tool only fills the labelled "PHOTO" placeholders it
finds, inside whatever frame that slide already has, and never touches
anything else. Slide numbers are the deck's own (as Google Slides shows them).

Requires: pip install python-pptx pillow
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

HERE = Path(__file__).resolve().parent.parent
DECK = HERE / "Project-One-Build-Walkthrough.pptx"
PAD, GUTTER = Inches(0.12), Inches(0.1)


def title_of(slide) -> str:
    texts = [s.text_frame.text.strip() for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
    return next((t for t in texts if len(t) < 70 and not t[:1].isdigit() and not t.startswith(("PHOTO", "DO", "CHECK", "OP "))), "")


def placeholder(slide):
    """The 'PHOTO …' text box, if the slide still has one."""
    for s in slide.shapes:
        if s.has_text_frame and s.text_frame.text.strip().startswith("PHOTO"):
            return s
    return None


def frame_around(slide, shape):
    """The empty frame shape that contains `shape`'s centre — the photo well."""
    cx, cy = shape.left + shape.width // 2, shape.top + shape.height // 2
    best = None
    for s in slide.shapes:
        if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and not (s.has_text_frame and s.text_frame.text.strip()):
            if s.left <= cx <= s.left + s.width and s.top <= cy <= s.top + s.height:
                if best is None or s.width * s.height < best.width * best.height:
                    best = s
    return best or shape


def pictures_in(slide, frame):
    out = []
    for s in slide.shapes:
        if s.shape_type == MSO_SHAPE_TYPE.PICTURE:
            cx, cy = s.left + s.width // 2, s.top + s.height // 2
            if frame.left <= cx <= frame.left + frame.width and frame.top <= cy <= frame.top + frame.height:
                out.append(s)
    return out


def layout(pics, frame):
    """Lay the pictures out side by side inside the frame, aspect preserved."""
    n = len(pics)
    x0, y0 = frame.left + PAD, frame.top + PAD
    w, h = frame.width - 2 * PAD - GUTTER * (n - 1), frame.height - 2 * PAD
    cw = w // n
    for i, p in enumerate(pics):
        iw, ih = p.image.size
        scale = min(cw / iw, h / ih)
        pw, ph = int(iw * scale), int(ih * scale)
        p.left = x0 + i * (cw + GUTTER) + (cw - pw) // 2
        p.top = y0 + (h - ph) // 2
        p.width, p.height = pw, ph


def list_slots(prs):
    for i, slide in enumerate(prs.slides, 1):
        ph = placeholder(slide)
        if ph:
            wanted = ph.text_frame.text.strip().split("\n")[-1].strip()
            print(f"slide {i:2d}  {title_of(slide)[:45]:<45}  waiting for: {wanted}")


def add(prs, n: int, photo: Path):
    slide = prs.slides[n - 1]
    ph = placeholder(slide)
    frame = frame_around(slide, ph) if ph else None
    if frame is None:
        # No placeholder left: reuse the frame around the existing pictures.
        pics = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        if not pics:
            sys.exit(f"slide {n} ({title_of(slide)}) has no PHOTO slot and no photos — nothing to add to")
        frame = frame_around(slide, pics[0])
    existing = pictures_in(slide, frame)
    pic = slide.shapes.add_picture(str(photo), frame.left + PAD, frame.top + PAD, width=Inches(1))
    if ph is not None:
        ph._element.getparent().remove(ph._element)
    layout(existing + [pic], frame)
    print(f"slide {n}  {title_of(slide)}: added {photo.name}")


def main(argv):
    if not DECK.exists():
        sys.exit(f"deck not found: {DECK}")
    prs = Presentation(DECK)
    if argv == ["--list"]:
        list_slots(prs)
        return
    if len(argv) != 2:
        sys.exit(__doc__)
    n, photo = int(argv[0]), (HERE / argv[1])
    if not photo.exists():
        sys.exit(f"photo not found: {photo}")
    add(prs, n, photo)
    prs.save(DECK)


if __name__ == "__main__":
    main(sys.argv[1:])
